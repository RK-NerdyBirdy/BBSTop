"""Read a CGPA sheet (CSV or .xlsx) into (register_number, raw cgpa) rows.

Columns are found by name, case-insensitively: anything containing "reg" for the register
number and "cgpa" for the CGPA, so official exports with extra columns work as-is.
"""

import csv
import io

from fastapi import HTTPException, status

MAX_ROWS = 5000


def _find_columns(header: list[str]) -> tuple[int, int] | None:
    names = [str(h or "").strip().lower().replace("_", " ").replace(".", " ") for h in header]
    reg = next((i for i, n in enumerate(names) if "reg" in n), None)
    cgpa = next((i for i, n in enumerate(names) if "cgpa" in n), None)
    return (reg, cgpa) if reg is not None and cgpa is not None else None


def _rows_from_table(table: list[list[object]]) -> list[tuple[int, str, str]]:
    """Returns (spreadsheet row number, register number, cgpa text)."""
    for header_index, header in enumerate(table[:10]):
        columns = _find_columns([str(c) if c is not None else "" for c in header])
        if columns:
            break
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Couldn't find the register number and CGPA columns. Name them e.g. 'register_number' and 'cgpa'.",
        )

    reg_col, cgpa_col = columns
    rows = []
    for offset, row in enumerate(table[header_index + 1:], start=header_index + 2):
        reg = str(row[reg_col]).strip() if reg_col < len(row) and row[reg_col] is not None else ""
        cgpa = str(row[cgpa_col]).strip() if cgpa_col < len(row) and row[cgpa_col] is not None else ""
        if not reg and not cgpa:
            continue
        rows.append((offset, reg.upper(), cgpa))
    if len(rows) > MAX_ROWS:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"File has more than {MAX_ROWS} rows")
    return rows


def read_cgpa_file(filename: str, content: bytes) -> list[tuple[int, str, str]]:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as error:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Couldn't read the Excel file") from error
        sheet = workbook.worksheets[0]
        table = [list(r) for r in sheet.iter_rows(values_only=True, max_row=MAX_ROWS + 20)]
        workbook.close()
        return _rows_from_table(table)
    if name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="CSV must be UTF-8 encoded") from error
        return _rows_from_table(list(csv.reader(io.StringIO(text))))
    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Upload a .csv or .xlsx file")
